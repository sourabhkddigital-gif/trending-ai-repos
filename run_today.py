"""Fetch today's GitHub trending repos, save them to a CSV, and print them.

Usage (from this folder):  python run_today.py
Output: daily/trending_today_YYYY-MM-DD.xlsx (clickable links) and .csv.
Does not change data/trending.csv or the README.
"""
import csv
import os
import re
import sys
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path

import tracker as t

sys.stdout.reconfigure(encoding="utf-8", errors="replace")


def today_ist() -> str:
    """Today's date in India time, so the GitHub server (UTC) and a local PC name files the same way."""
    return datetime.now(timezone(timedelta(hours=5, minutes=30))).date().isoformat()


def clean(repo: str, desc: str) -> str:
    # Remove the "Sponsor Star owner / name" button text the trending page leaks into descriptions.
    return re.sub(r"^(?:Sponsor )?Star " + re.escape(repo.replace("/", " / ")) + r" ?", "", desc).strip()


def main() -> int:
    print("Fetching today's trending repositories from github.com/trending ...")
    repos, source = t.fetch_trending()
    if not repos:
        print("Could not reach GitHub. Check your internet connection and try again.")
        return 1
    print(f"  {len(repos)} repos found ({source}). Looking up topics ...")

    rows = []
    for i, r in enumerate(repos, 1):
        r["description"] = clean(r["repo"], r["description"])
        t.enrich(r)
        rows.append({
            "rank": i,
            "repo": r["repo"],
            "url": f"https://github.com/{r['repo']}",
            "description": r["description"],
            "language": r["language"],
            "stars_today": r["stars_today"],
            "total_stars": r["stars"],
            "forks": r["forks"],
            "topics": r["topics"],
            "license": r["license"],
            "is_ai": "yes" if t.is_ai(r) else "no",
        })
        time.sleep(0.3)

    folder = Path(__file__).resolve().parent / "daily"
    folder.mkdir(exist_ok=True)
    out = folder / f"trending_today_{today_ist()}.csv"
    try:
        with out.open("w", encoding="utf-8-sig", newline="") as f:  # utf-8-sig so Excel shows symbols correctly
            w = csv.DictWriter(f, fieldnames=list(rows[0]))
            w.writeheader()
            w.writerows(rows)
    except PermissionError:
        print(f"Could not save {out.name}: close it in Excel and run again.")
        return 1
    xlsx = save_excel(rows, out.with_suffix(".xlsx"))

    print()
    print(f"{'#':>2}  {'Repository':<38} {'Language':<11} {'Stars today':>11} {'Total':>9}  AI")
    print("-" * 80)
    for r in rows:
        print(f"{r['rank']:>2}  {r['repo'][:38]:<38} {(r['language'] or '-')[:11]:<11} "
              f"{r['stars_today']:>11,} {r['total_stars']:>9,}  {r['is_ai']}")
    print("-" * 80)
    print(f"{sum(r['is_ai'] == 'yes' for r in rows)} of {len(rows)} are AI-related.")
    print(f"Saved: {out}")
    if xlsx:
        print(f"Saved: {xlsx}  (open this one in Excel for clickable links)")
    # On a Windows PC, open today's file straight away. Skipped on GitHub's servers (CI is set there).
    if sys.platform == "win32" and not os.getenv("CI"):
        print("Opening it in Excel ...")
        os.startfile(xlsx or out)
    return 0


def save_excel(rows: list[dict], path: Path) -> Path | None:
    """Write a real Excel file with clickable repo links. Skipped if openpyxl is missing."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        print("  (Install openpyxl for an Excel file with clickable links: pip install openpyxl)")
        return None
    wb = Workbook()
    ws = wb.active
    ws.title = "Trending today"
    headers = ["Rank", "Repository", "Link", "Description", "Language", "Stars today",
               "Total stars", "Forks", "Topics", "License", "AI"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1D7A4C")
    link_font = Font(color="0563C1", underline="single")
    for r in rows:
        ws.append([r["rank"], r["repo"], r["url"], r["description"], r["language"], r["stars_today"],
                   r["total_stars"], r["forks"], r["topics"], r["license"], r["is_ai"]])
        row = ws.max_row
        for col in (2, 3):  # repo name and link both open the repo
            cell = ws.cell(row=row, column=col)
            cell.hyperlink = r["url"]
            cell.font = link_font
        for col in (6, 7, 8):
            ws.cell(row=row, column=col).number_format = "#,##0"
    for col, width in zip("ABCDEFGHIJK", (6, 34, 48, 70, 13, 12, 12, 10, 40, 14, 6)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    try:
        wb.save(path)
    except PermissionError:
        print(f"  Could not save {path.name}: close it in Excel and run again.")
        return None
    return path


if __name__ == "__main__":
    sys.exit(main())
